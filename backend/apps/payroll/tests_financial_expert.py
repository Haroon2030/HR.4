"""
فحص مالي لمحرك المسير — اختبارات متنوعة من منظور محاسبي.

كل اختبار يصف القاعدة المالية المتوقعة صراحةً. الاختبارات التي تفشل هنا تكشف
انحرافاً عن القاعدة (مع تعليق يشرح الأثر المالي).
"""
from datetime import date
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase

from apps.core.models import Branch, Company
from apps.employees.models import (
    Employee,
    EmployeeAbsence,
    EmployeeCashShortage,
    EmployeeLeave,
    EmployeeLedger,
    EmployeeLoan,
    EmployeeStatement,
    LoanInstallment,
)
from apps.payroll.models import PayrollLine, PayrollRun
from apps.payroll.services.engine import (
    build_consolidated_payroll_run,
    build_payroll_run,
    lock_payroll_run,
    unlock_payroll_run,
)
from apps.payroll.services.financial_audit import audit_payroll_runs
from apps.setup.models import Sponsorship

User = get_user_model()
import unittest
# عيب مالي مؤكَّد في المحرك — يُزال الوسم عند الإصلاح (سيفشل الاختبار «نجح بشكل غير متوقع» تنبيهاً)
known_defect = unittest.expectedFailure
D = Decimal
TRANSFER = PayrollRun.SalaryMode.TRANSFER
CASH = PayrollRun.SalaryMode.CASH


class FinanceBase(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name='شركة الفحص')
        self.branch = Branch.objects.create(name='فرع 1', code='F1', company=self.company)
        self.branch2 = Branch.objects.create(name='فرع 2', code='F2', company=self.company)
        self.user = User.objects.create_user(username='fin_auditor', password='x-pass-123')
        self.sp = Sponsorship.objects.create(code='SPF', company_name='كفالة الفحص')

    def emp(self, name='موظف', *, branch=None, sponsored=True, basic='3000', housing='1000',
            transport='500', other='0', meal='0', rate='10', hire=date(2020, 1, 1),
            end=None, status=Employee.Status.ACTIVE):
        return Employee.objects.create(
            name=name, branch=branch or self.branch,
            sponsorship=self.sp if sponsored else None,
            status=status, hire_date=hire, end_date=end,
            basic_salary=D(basic), housing_allowance=D(housing),
            transport_allowance=D(transport), other_allowance=D(other),
            meal_allowance=D(meal), insurance_deduction_rate=D(rate),
        )

    def build(self, year, month, *, sponsored=True, branch=None):
        return build_payroll_run(
            branch or self.branch, year, month, self.user,
            salary_mode=TRANSFER if sponsored else CASH,
            sponsorship_id=self.sp.id if sponsored else None,
        )

    def line(self, run, emp):
        return run.lines.get(employee=emp)


# ══════════════════════════════════════════════════════════════════════════════
# 1) الاستحقاق الجزئي للشهر (تاريخ المباشرة/التوقف) وأطوال الأشهر
# ══════════════════════════════════════════════════════════════════════════════
class ProrationTests(FinanceBase):
    def test_full_month_of_30_days_pays_full_gross(self):
        e = self.emp(rate='0')
        self.assertEqual(self.line(self.build(2026, 4), e).gross_salary, D('4500.00'))

    def test_full_month_of_31_days_pays_full_gross(self):
        e = self.emp(rate='0')
        self.assertEqual(self.line(self.build(2026, 3), e).gross_salary, D('4500.00'))

    def test_full_february_pays_full_gross(self):
        """موظف على رأس العمل طوال فبراير يجب أن يستلم راتبه كاملاً (شهر = 30 يوماً)."""
        e = self.emp(rate='0')
        self.assertEqual(self.line(self.build(2026, 2), e).gross_salary, D('4500.00'))

    def test_full_leap_february_pays_full_gross(self):
        e = self.emp(rate='0')
        self.assertEqual(self.line(self.build(2028, 2), e).gross_salary, D('4500.00'))

    def test_mid_month_hire_is_prorated(self):
        e = self.emp(rate='0', hire=date(2026, 4, 16))  # 15 يوماً من 30
        self.assertEqual(self.line(self.build(2026, 4), e).gross_salary, D('2250.00'))

    def test_mid_month_hire_in_31_day_month_follows_30_day_convention(self):
        """مباشرة 16/3: شهر محاسبي 30 يوماً ⇒ 15 يوماً مستحقة (لا 16)."""
        e = self.emp(rate='0', hire=date(2026, 3, 16))
        self.assertEqual(self.line(self.build(2026, 3), e).gross_salary, D('2250.00'))

    def test_mid_month_hire_in_february_follows_30_day_convention(self):
        """مباشرة 16/2: 15 يوماً مستحقة بقاعدة الـ 30 يوماً (لا 13)."""
        e = self.emp(rate='0', hire=date(2026, 2, 16))
        self.assertEqual(self.line(self.build(2026, 2), e).gross_salary, D('2250.00'))

    def test_end_date_on_last_day_of_31_day_month_pays_full(self):
        e = self.emp(rate='0', end=date(2026, 3, 31))
        self.assertEqual(self.line(self.build(2026, 3), e).gross_salary, D('4500.00'))

    def test_first_day_hire_pays_full(self):
        e = self.emp(rate='0', hire=date(2026, 4, 1))
        self.assertEqual(self.line(self.build(2026, 4), e).gross_salary, D('4500.00'))

    def test_hire_after_period_is_excluded(self):
        e = self.emp(hire=date(2026, 5, 1))
        run = self.build(2026, 4)
        self.assertFalse(run.lines.filter(employee=e).exists())

    def test_end_date_mid_month_is_prorated(self):
        e = self.emp(rate='0', end=date(2026, 4, 15))
        self.assertEqual(self.line(self.build(2026, 4), e).gross_salary, D('2250.00'))

    def test_ended_before_period_is_excluded(self):
        e = self.emp(end=date(2026, 3, 31))
        self.assertFalse(self.build(2026, 4).lines.filter(employee=e).exists())

    def test_hire_and_end_in_same_month(self):
        e = self.emp(rate='0', hire=date(2026, 4, 10), end=date(2026, 4, 19))  # 10 أيام
        self.assertEqual(self.line(self.build(2026, 4), e).gross_salary, D('1500.00'))


# ══════════════════════════════════════════════════════════════════════════════
# 2) الغياب والإجازة بدون راتب
# ══════════════════════════════════════════════════════════════════════════════
class AbsenceAndLeaveTests(FinanceBase):
    def test_single_absence_day_is_gross_over_30(self):
        e = self.emp(rate='0')
        EmployeeAbsence.objects.create(employee=e, absence_date=date(2026, 4, 5), days=1)
        ln = self.line(self.build(2026, 4), e)
        self.assertEqual(ln.absence_deduction, D('150.00'))
        self.assertEqual(ln.net_salary, D('4350.00'))

    def test_multiple_absence_records_accumulate(self):
        e = self.emp(rate='0')
        EmployeeAbsence.objects.create(employee=e, absence_date=date(2026, 4, 5), days=1)
        EmployeeAbsence.objects.create(employee=e, absence_date=date(2026, 4, 9), days=2)
        ln = self.line(self.build(2026, 4), e)
        self.assertEqual(ln.absence_deduction, D('450.00'))
        self.assertEqual(ln.absence_days, D('3'))

    def test_absence_outside_period_is_ignored(self):
        e = self.emp(rate='0')
        EmployeeAbsence.objects.create(employee=e, absence_date=date(2026, 3, 31), days=1)
        EmployeeAbsence.objects.create(employee=e, absence_date=date(2026, 5, 1), days=1)
        self.assertEqual(self.line(self.build(2026, 4), e).absence_deduction, D('0.00'))

    def test_full_month_absence_zeroes_the_gross_exactly(self):
        """غياب 30 يوماً = خصم الراتب كاملاً بدون فروق تقريب."""
        e = self.emp(basic='1000', housing='0', transport='0', rate='0')
        EmployeeAbsence.objects.create(employee=e, absence_date=date(2026, 4, 1), days=30)
        ln = self.line(self.build(2026, 4), e)
        self.assertEqual(ln.absence_deduction, D('1000.00'))
        self.assertEqual(ln.net_salary, D('0.00'))

    def test_unpaid_leave_counts_only_days_inside_period(self):
        e = self.emp(rate='0')
        EmployeeLeave.objects.create(
            employee=e, leave_type=EmployeeLeave.LeaveType.UNPAID,
            date_from=date(2026, 3, 28), date_to=date(2026, 4, 3), days=D('7'),
        )
        ln = self.line(self.build(2026, 4), e)
        self.assertEqual(ln.unpaid_leave_days, D('3'))
        self.assertEqual(ln.unpaid_leave_deduction, D('450.00'))

    def test_unpaid_leave_for_whole_31_day_month_never_exceeds_gross(self):
        """إجازة بدون راتب لشهر 31 يوماً: الخصم لا يتجاوز الراتب (الشهر المحاسبي 30 يوماً)."""
        e = self.emp(rate='0')
        EmployeeLeave.objects.create(
            employee=e, leave_type=EmployeeLeave.LeaveType.UNPAID,
            date_from=date(2026, 3, 1), date_to=date(2026, 3, 31), days=D('31'),
        )
        ln = self.line(self.build(2026, 3), e)
        self.assertLessEqual(ln.unpaid_leave_deduction, D('4500.00'))
        self.assertEqual(Decimal(ln.breakdown['uncollected_deductions']), D('0.00'))

    def test_absence_inside_unpaid_leave_is_not_deducted_twice(self):
        e = self.emp(rate='0')
        EmployeeLeave.objects.create(
            employee=e, leave_type=EmployeeLeave.LeaveType.UNPAID,
            date_from=date(2026, 4, 10), date_to=date(2026, 4, 12), days=D('3'),
        )
        EmployeeAbsence.objects.create(employee=e, absence_date=date(2026, 4, 11), days=1)
        ln = self.line(self.build(2026, 4), e)
        self.assertEqual(ln.absence_deduction + ln.unpaid_leave_deduction, D('450.00'))

    def test_paid_leave_types_do_not_deduct(self):
        e = self.emp(rate='0')
        EmployeeLeave.objects.create(
            employee=e, leave_type=EmployeeLeave.LeaveType.ANNUAL,
            date_from=date(2026, 4, 1), date_to=date(2026, 4, 10), days=D('10'),
        )
        self.assertEqual(self.line(self.build(2026, 4), e).unpaid_leave_deduction, D('0.00'))

    def test_deductions_above_salary_are_capped_and_flagged(self):
        e = self.emp(basic='1000', housing='0', transport='0', rate='0')
        EmployeeStatement.objects.create(
            employee=e, statement_type=EmployeeStatement.StatementType.PENALTY,
            title='غرامة', statement_date=date(2026, 4, 3), deduction_amount=D('1800'),
        )
        ln = self.line(self.build(2026, 4), e)
        self.assertEqual(ln.net_salary, D('0.00'))               # لا صافي سالب
        self.assertEqual(Decimal(ln.breakdown['uncollected_deductions']), D('800.00'))


# ══════════════════════════════════════════════════════════════════════════════
# 3) التأمينات
# ══════════════════════════════════════════════════════════════════════════════
class InsuranceTests(FinanceBase):
    def test_insurance_base_is_basic_plus_housing_only(self):
        e = self.emp(rate='10')  # (3000+1000)×10% = 400 — النقل مستثنى
        self.assertEqual(self.line(self.build(2026, 4), e).insurance_deduction, D('400.00'))

    def test_insurance_rate_is_clamped_between_0_and_100(self):
        hi = self.emp('أ', rate='250')
        lo = self.emp('ب', rate='-5')
        run = self.build(2026, 4)
        self.assertEqual(self.line(run, hi).insurance_deduction, D('4000.00'))
        self.assertEqual(self.line(run, lo).insurance_deduction, D('0.00'))

    def test_insurance_is_prorated_with_salary(self):
        e = self.emp(rate='10', hire=date(2026, 4, 16))
        self.assertEqual(self.line(self.build(2026, 4), e).insurance_deduction, D('200.00'))

    def test_meal_allowance_not_in_insurance_base(self):
        e = self.emp(rate='10', meal='600')
        ln = self.line(self.build(2026, 4), e)
        self.assertEqual(ln.insurance_deduction, D('400.00'))
        self.assertEqual(ln.gross_salary, D('5100.00'))


# ══════════════════════════════════════════════════════════════════════════════
# 4) التقريب والتطابق الحسابي
# ══════════════════════════════════════════════════════════════════════════════
class RoundingAndReconciliationTests(FinanceBase):
    def test_every_line_satisfies_net_equals_earnings_minus_deductions(self):
        for i, basic in enumerate(['3333.33', '2750.55', '1999.99', '4100.10', '999.95']):
            e = self.emp(f'م{i}', basic=basic, housing='1111.11', transport='222.22', rate='9.75')
            EmployeeAbsence.objects.create(employee=e, absence_date=date(2026, 4, 3), days=i + 1)
        run = self.build(2026, 4)
        for ln in run.lines.all():
            self.assertEqual(ln.net_salary, ln.total_earnings - ln.total_deductions)
            for f in ('gross_salary', 'net_salary', 'total_deductions', 'insurance_deduction',
                      'absence_deduction'):
                v = getattr(ln, f)
                self.assertEqual(v, v.quantize(D('0.01')), f)
            self.assertGreaterEqual(ln.net_salary, 0)

    def test_run_totals_equal_sum_of_lines(self):
        for i in range(7):
            self.emp(f'م{i}', basic=str(2000 + i * 137.37), rate='9')
        run = self.build(2026, 4)
        agg = run.lines.aggregate_sums = None
        from django.db.models import Sum
        s = run.lines.aggregate(e=Sum('total_earnings'), d=Sum('total_deductions'), n=Sum('net_salary'))
        run.refresh_from_db()
        self.assertEqual((run.total_earnings, run.total_deductions, run.total_net),
                         (s['e'], s['d'], s['n']))
        self.assertEqual(run.employees_count, 7)
        self.assertEqual(run.total_net, run.total_earnings - run.total_deductions)

    def test_rebuild_is_idempotent(self):
        e = self.emp()
        EmployeeAbsence.objects.create(employee=e, absence_date=date(2026, 4, 5), days=2)
        first = self.build(2026, 4)
        snap = (first.total_net, first.employees_count, first.lines.count())
        again = self.build(2026, 4)
        self.assertEqual(first.pk, again.pk)
        self.assertEqual((again.total_net, again.employees_count, again.lines.count()), snap)

    def test_zero_salary_employee_does_not_break_the_run(self):
        e = self.emp(basic='0', housing='0', transport='0', rate='10')
        run = self.build(2026, 4)
        ln = self.line(run, e)
        self.assertEqual((ln.gross_salary, ln.net_salary), (D('0.00'), D('0.00')))
        lock_payroll_run(run, self.user)  # يُرحَّل بلا خطأ

    def test_negative_net_is_impossible(self):
        e = self.emp(basic='500', housing='0', transport='0', rate='100')
        EmployeeStatement.objects.create(
            employee=e, statement_type=EmployeeStatement.StatementType.PENALTY,
            title='غرامة', statement_date=date(2026, 4, 3), deduction_amount=D('900'),
        )
        EmployeeCashShortage.objects.create(
            employee=e, branch=self.branch, shortage_date=date(2026, 4, 4), amount=D('700'),
        )
        self.assertGreaterEqual(self.line(self.build(2026, 4), e).net_salary, 0)


# ══════════════════════════════════════════════════════════════════════════════
# 5) السلف
# ══════════════════════════════════════════════════════════════════════════════
class LoanTests(FinanceBase):
    def _loan(self, e, amount, installments, monthly, first=date(2026, 4, 1)):
        loan = EmployeeLoan.objects.create(
            employee=e, amount=D(amount), monthly_deduction=D(monthly),
            installments=installments, issued_at=date(2026, 3, 1), first_deduction_date=first,
        )
        loan.generate_installments()
        return loan

    def test_installment_reduces_net_and_is_marked_paid_on_lock(self):
        e = self.emp(rate='0')
        loan = self._loan(e, '1000', 4, '250')
        run = self.build(2026, 4)
        self.assertEqual(self.line(run, e).loan_deduction, D('250.00'))
        self.assertEqual(self.line(run, e).net_salary, D('4250.00'))
        lock_payroll_run(run, self.user)
        self.assertEqual(loan.installments_log.get(period_month=4).status, LoanInstallment.Status.PAID)
        self.assertEqual(loan.installments_log.filter(status=LoanInstallment.Status.PENDING).count(), 3)

    def test_last_installment_closes_the_loan_and_unlock_reopens_it(self):
        e = self.emp(rate='0')
        loan = self._loan(e, '250', 1, '250')
        run = self.build(2026, 4)
        lock_payroll_run(run, self.user)
        loan.refresh_from_db()
        self.assertEqual(loan.status, EmployeeLoan.Status.PAID)
        unlock_payroll_run(run, self.user)
        loan.refresh_from_db()
        self.assertEqual(loan.status, EmployeeLoan.Status.ACTIVE)
        self.assertEqual(loan.installments_log.get().status, LoanInstallment.Status.PENDING)

    def test_two_loans_deducted_together(self):
        e = self.emp(rate='0')
        self._loan(e, '600', 3, '200')
        self._loan(e, '900', 3, '300')
        self.assertEqual(self.line(self.build(2026, 4), e).loan_deduction, D('500.00'))

    def test_installments_never_push_net_below_zero(self):
        e = self.emp(basic='1000', housing='0', transport='0', rate='0')
        self._loan(e, '6000', 3, '2000')
        ln = self.line(self.build(2026, 4), e)
        self.assertEqual(ln.loan_deduction, D('0.00'))
        self.assertEqual(ln.net_salary, D('1000.00'))
        self.assertEqual(len(ln.breakdown['deferred_installments']), 1)

    def test_installment_of_other_month_is_not_taken(self):
        e = self.emp(rate='0')
        self._loan(e, '1000', 4, '250', first=date(2026, 6, 1))
        self.assertEqual(self.line(self.build(2026, 4), e).loan_deduction, D('0.00'))


# ══════════════════════════════════════════════════════════════════════════════
# 6) الترحيل: لا ازدواج، ذرّية، قيود المخصصات
# ══════════════════════════════════════════════════════════════════════════════
class LockIntegrityTests(FinanceBase):
    def test_lock_twice_is_refused(self):
        self.emp()
        run = self.build(2026, 4)
        lock_payroll_run(run, self.user)
        with self.assertRaises(ValueError):
            lock_payroll_run(run, self.user)
        self.assertEqual(EmployeeLedger.objects.filter(payroll_run=run).count(), 1)

    def test_locked_run_cannot_be_rebuilt(self):
        self.emp()
        run = self.build(2026, 4)
        lock_payroll_run(run, self.user)
        with self.assertRaises(ValueError):
            self.build(2026, 4)

    def test_deduction_items_are_bound_once_and_released_on_unlock(self):
        e = self.emp(rate='0')
        ab = EmployeeAbsence.objects.create(employee=e, absence_date=date(2026, 4, 5), days=1)
        run = self.build(2026, 4)
        lock_payroll_run(run, self.user)
        ab.refresh_from_db()
        self.assertEqual(ab.applied_to_payroll_id, run.pk)
        unlock_payroll_run(run, self.user)
        ab.refresh_from_db()
        self.assertIsNone(ab.applied_to_payroll_id)
        # بعد الفتح يظهر الغياب في إعادة البناء ثم يُربط مرة واحدة فقط
        run = self.build(2026, 4)
        self.assertEqual(self.line(run, e).absence_deduction, D('150.00'))

    def test_failed_lock_leaves_every_item_unbound(self):
        from unittest import mock
        e = self.emp(rate='0')
        ab = EmployeeAbsence.objects.create(employee=e, absence_date=date(2026, 4, 5), days=1)
        run = self.build(2026, 4)
        with mock.patch(
            'apps.employees.services.accrual_ledger_notes.compute_monthly_ledger_amounts',
            side_effect=RuntimeError('boom'),
        ):
            with self.assertRaises(ValueError):
                lock_payroll_run(run, self.user)
        ab.refresh_from_db()
        self.assertIsNone(ab.applied_to_payroll_id)
        run.refresh_from_db()
        self.assertEqual(run.status, PayrollRun.Status.DRAFT)
        self.assertEqual(EmployeeLedger.objects.filter(payroll_run=run).count(), 0)

    def test_same_employee_cannot_be_paid_twice_for_same_month(self):
        """موظف رُحِّل في مسير موحّد لا يجوز أن يدخل مسير فرع قياسي لنفس الشهر."""
        e = self.emp(rate='0')
        cons = build_consolidated_payroll_run(
            [self.branch], 2026, 4, self.user, salary_mode=TRANSFER, sponsorship_id=self.sp.id,
        )
        lock_payroll_run(cons, self.user)
        std = self.build(2026, 4)
        self.assertFalse(std.lines.filter(employee=e).exists(),
                         'الموظف مُرحَّل في المسير الموحّد ويظهر مرة ثانية في مسير الفرع')

    def test_standard_locked_then_consolidated_skips_employee(self):
        e = self.emp(rate='0')
        std = self.build(2026, 4)
        lock_payroll_run(std, self.user)
        cons = build_consolidated_payroll_run(
            [self.branch], 2026, 4, self.user, salary_mode=TRANSFER, sponsorship_id=self.sp.id,
        )
        self.assertTrue(cons is None or not cons.lines.filter(employee=e).exists())

    def test_two_overlapping_drafts_cannot_both_be_locked(self):
        """مسودتان تضمّان الموظف نفسه: الترحيل الثاني يجب أن يُرفض (منع صرفه مرتين)."""
        e = self.emp(rate='0')
        cons = build_consolidated_payroll_run(
            [self.branch], 2026, 4, self.user, salary_mode=TRANSFER, sponsorship_id=self.sp.id,
        )
        std = self.build(2026, 4)          # مسودة قياسية تُبنى بعد الموحّدة — الاثنتان تضمّان الموظف
        self.assertTrue(cons.lines.filter(employee=e).exists())
        self.assertTrue(std.lines.filter(employee=e).exists())
        lock_payroll_run(cons, self.user)
        with self.assertRaises(ValueError):
            lock_payroll_run(std, self.user)
        self.assertEqual(EmployeeLedger.objects.filter(employee=e, date=date(2026, 4, 30)).count(), 1)


class AccrualLedgerTests(FinanceBase):
    def test_monthly_ledger_amounts_for_sponsored_employee(self):
        e = self.emp(rate='0', hire=date(2023, 1, 1))  # إجمالي 4500 — خدمة < 5 سنوات
        run = self.build(2026, 4)
        lock_payroll_run(run, self.user)
        lg = EmployeeLedger.objects.get(payroll_run=run, employee=e)
        self.assertEqual(lg.date, date(2026, 4, 30))
        self.assertEqual(lg.leave_days_change, D('1.75'))
        self.assertEqual(lg.leave_amount_change, D('262.50'))      # 1.75 × (4500/30)
        self.assertEqual(lg.eosb_amount_change, D('187.50'))       # 4500 ÷ 24 (≤ 5 سنوات)

    def test_eosb_after_five_years_is_gross_over_12(self):
        e = self.emp(rate='0', hire=date(2018, 1, 1))
        run = self.build(2026, 4)
        lock_payroll_run(run, self.user)
        self.assertEqual(
            EmployeeLedger.objects.get(payroll_run=run, employee=e).eosb_amount_change, D('375.00'),
        )

    def test_eosb_base_excludes_meal_allowance(self):
        e = self.emp(rate='0', meal='600', hire=date(2023, 1, 1))  # إجمالي 5100، الأساس 4500
        run = self.build(2026, 4)
        lock_payroll_run(run, self.user)
        self.assertEqual(
            EmployeeLedger.objects.get(payroll_run=run, employee=e).eosb_amount_change, D('187.50'),
        )

    def test_cash_employee_accrues_leave_but_no_eosb(self):
        e = self.emp(sponsored=False, rate='0')
        run = self.build(2026, 4, sponsored=False)
        lock_payroll_run(run, self.user)
        lg = EmployeeLedger.objects.get(payroll_run=run, employee=e)
        self.assertEqual(lg.eosb_amount_change, D('0'))
        self.assertEqual(lg.leave_days_change, D('1.75'))

    def test_cumulative_balances_chain_across_months(self):
        e = self.emp(rate='0')
        m4 = self.build(2026, 4)
        lock_payroll_run(m4, self.user)
        m5 = self.build(2026, 5)
        lock_payroll_run(m5, self.user)
        l4 = EmployeeLedger.objects.get(payroll_run=m4, employee=e)
        l5 = EmployeeLedger.objects.get(payroll_run=m5, employee=e)
        self.assertEqual(l5.cumulative_eosb_amount, l4.cumulative_eosb_amount + l5.eosb_amount_change)
        self.assertEqual(l5.cumulative_leave_days, l4.cumulative_leave_days + D('1.75'))

    def test_unlock_removes_ledger_and_relock_recreates_exactly_one(self):
        e = self.emp(rate='0')
        run = self.build(2026, 4)
        lock_payroll_run(run, self.user)
        unlock_payroll_run(run, self.user)
        self.assertEqual(EmployeeLedger.objects.filter(payroll_run=run).count(), 0)
        run = self.build(2026, 4)
        lock_payroll_run(run, self.user)
        self.assertEqual(EmployeeLedger.objects.filter(payroll_run=run, employee=e).count(), 1)

    def test_mid_month_hire_accrues_eosb_on_prorated_gross(self):
        """الاستحقاق الشهري يُحسب على الراتب المستحق فعلياً (جزئي) لا الكامل."""
        e = self.emp(rate='0', hire=date(2026, 4, 16))
        run = self.build(2026, 4)
        lock_payroll_run(run, self.user)
        self.assertEqual(
            EmployeeLedger.objects.get(payroll_run=run, employee=e).eosb_amount_change,
            D('93.75'),  # 2250 ÷ 24
        )


# ══════════════════════════════════════════════════════════════════════════════
# 7) فصل الأنماط والحالات + التوحيد + التدقيق
# ══════════════════════════════════════════════════════════════════════════════
class ScopeAndAuditTests(FinanceBase):
    def test_cash_and_transfer_runs_are_disjoint(self):
        s = self.emp('بكفالة', sponsored=True)
        c = self.emp('بدون كفالة', sponsored=False)
        t = self.build(2026, 4, sponsored=True)
        k = self.build(2026, 4, sponsored=False)
        self.assertEqual(list(t.lines.values_list('employee_id', flat=True)), [s.id])
        self.assertEqual(list(k.lines.values_list('employee_id', flat=True)), [c.id])

    def test_only_active_and_on_leave_employees_are_paid(self):
        a = self.emp('نشط')
        l = self.emp('إجازة', status=Employee.Status.LEAVE)
        self.emp('موقوف', status=Employee.Status.SUSPENDED)
        self.emp('منتهي', status=Employee.Status.TERMINATED)
        ids = set(self.build(2026, 4).lines.values_list('employee_id', flat=True))
        self.assertEqual(ids, {a.id, l.id})

    def test_cash_net_is_all_cash_and_transfer_net_is_all_bank(self):
        s = self.emp('بكفالة', sponsored=True)
        c = self.emp('نقدي', sponsored=False)
        ls = self.line(self.build(2026, 4, sponsored=True), s)
        lc = self.line(self.build(2026, 4, sponsored=False), c)
        self.assertEqual((ls.net_cash_amount, ls.net_bank_transfer), (D('0'), ls.net_salary))
        self.assertEqual((lc.net_cash_amount, lc.net_bank_transfer), (lc.net_salary, D('0')))

    def test_consolidated_equals_sum_of_branch_runs(self):
        for i in range(3):
            self.emp(f'ف1-{i}', branch=self.branch, basic=str(2500 + i * 111))
            self.emp(f'ف2-{i}', branch=self.branch2, basic=str(2600 + i * 77))
        a = self.build(2026, 4, branch=self.branch)
        b = self.build(2026, 4, branch=self.branch2)
        expected = (a.total_net + b.total_net, a.total_earnings + b.total_earnings,
                    a.total_deductions + b.total_deductions, a.employees_count + b.employees_count)
        cons = build_consolidated_payroll_run(
            [self.branch, self.branch2], 2026, 4, self.user,
            salary_mode=TRANSFER, sponsorship_id=self.sp.id,
        )
        self.assertEqual(
            (cons.total_net, cons.total_earnings, cons.total_deductions, cons.employees_count),
            (expected[0], expected[1], expected[2], expected[3]),
        )

    def test_audit_is_clean_for_a_consistent_run(self):
        e = self.emp()
        EmployeeAbsence.objects.create(employee=e, absence_date=date(2026, 4, 5), days=1)
        run = self.build(2026, 4)
        audit = audit_payroll_runs([run])
        self.assertEqual(audit.error_count, 0, [c.detail for c in audit.checks if c.level == 'error'])
        self.assertTrue(audit.ready_to_lock)

    def test_audit_detects_tampered_line_and_blocks_readiness(self):
        e = self.emp()
        run = self.build(2026, 4)
        PayrollLine.objects.filter(run=run, employee=e).update(net_salary=D('9999.00'))
        audit = audit_payroll_runs([run])
        self.assertGreater(audit.error_count, 0)
        self.assertFalse(audit.ready_to_lock)

    def test_audit_detects_header_total_mismatch(self):
        self.emp()
        run = self.build(2026, 4)
        PayrollRun.objects.filter(pk=run.pk).update(total_net=D('1.00'))
        run.refresh_from_db()
        self.assertGreater(audit_payroll_runs([run]).error_count, 0)

    def test_audit_detects_wrong_insurance(self):
        e = self.emp(rate='10')
        run = self.build(2026, 4)
        PayrollLine.objects.filter(run=run, employee=e).update(insurance_deduction=D('1.00'))
        self.assertGreater(audit_payroll_runs([run]).error_count, 0)


# ══════════════════════════════════════════════════════════════════════════════
# 8) سيناريوهات تشغيلية أوسع (فحص المدير المالي — الجولة الثانية)
# ══════════════════════════════════════════════════════════════════════════════
class OperationalScenarioTests(FinanceBase):
    def _loan(self, e, amount='1000', n=4, monthly='250'):
        loan = EmployeeLoan.objects.create(
            employee=e, amount=D(amount), monthly_deduction=D(monthly), installments=n,
            issued_at=date(2026, 3, 1), first_deduction_date=date(2026, 4, 1),
        )
        loan.generate_installments()
        return loan

    def test_deleted_loan_is_not_deducted(self):
        """سلفة حُذفت من ملف الموظف: أقساطها يجب ألا تُخصم من الراتب."""
        e = self.emp(rate='0')
        self._loan(e).delete()
        self.assertEqual(self.line(self.build(2026, 4), e).loan_deduction, D('0.00'))

    def test_future_month_cannot_be_locked(self):
        """ترحيل شهر لم يبدأ بعد (مثلاً 2030) يجب أن يُرفض."""
        self.emp(rate='0')
        run = self.build(2030, 1)
        with self.assertRaises(ValueError):
            lock_payroll_run(run, self.user)

    def test_excel_export_totals_match_run_totals(self):
        from apps.payroll.services.export_excel import PAYROLL_EXPORT_COLUMNS, build_payroll_run_workbook
        for i in range(4):
            e = self.emp(f'م{i}', basic=str(2100 + i * 333.33), rate='9.75')
            EmployeeAbsence.objects.create(employee=e, absence_date=date(2026, 4, 2), days=i)
        run = self.build(2026, 4)
        run.refresh_from_db()
        ws = build_payroll_run_workbook(run).active
        col = next(i for i, c in enumerate(PAYROLL_EXPORT_COLUMNS, start=1) if c[0] == 'net_salary')
        footer = 2 + run.employees_count
        self.assertEqual(ws.cell(row=footer, column=1 if False else col).value is not None, True)
        self.assertAlmostEqual(float(ws.cell(row=footer, column=col).value), float(run.total_net), places=2)

    def test_user_without_manage_permission_cannot_build_or_lock(self):
        from django.urls import reverse
        self.emp(rate='0')
        viewer = User.objects.create_user(username='viewer_only', password='x-pass-123')
        self.client.login(username='viewer_only', password='x-pass-123')
        self.client.post(reverse('web:list_payroll_runs'),
                         {'year': 2026, 'month': 4, 'salary_mode': 'transfer', 'build_kind': 'standard'})
        self.assertFalse(PayrollRun.objects.exists())


# ══════════════════════════════════════════════════════════════════════════════
# 8) سيناريوهات تشغيلية متقدمة (ما يحدث في الواقع بعد الترحيل / بأثر رجعي)
# ══════════════════════════════════════════════════════════════════════════════
class AdvancedOperationalScenarioTests(FinanceBase):
    # ── بنود تُسجَّل بأثر رجعي بعد ترحيل شهرها ─────────────────────────────
    def test_backdated_absence_after_lock_is_recovered_next_month(self):
        """غياب أبريل سُجِّل بعد ترحيل أبريل: يجب أن يُخصم في مايو لا أن يضيع."""
        e = self.emp(rate='0')
        lock_payroll_run(self.build(2026, 4), self.user)
        EmployeeAbsence.objects.create(employee=e, absence_date=date(2026, 4, 20), days=1)
        self.assertEqual(self.line(self.build(2026, 5), e).absence_deduction, D('150.00'))

    def test_backdated_penalty_after_lock_is_recovered_next_month(self):
        e = self.emp(rate='0')
        lock_payroll_run(self.build(2026, 4), self.user)
        EmployeeStatement.objects.create(
            employee=e, statement_type=EmployeeStatement.StatementType.PENALTY,
            title='مخالفة أبريل', statement_date=date(2026, 4, 25), deduction_amount=D('200'),
        )
        self.assertEqual(self.line(self.build(2026, 5), e).penalty_deduction, D('200.00'))

    def test_backdated_cash_shortage_after_lock_is_recovered_next_month(self):
        e = self.emp(rate='0')
        lock_payroll_run(self.build(2026, 4), self.user)
        EmployeeCashShortage.objects.create(
            employee=e, branch=self.branch, shortage_date=date(2026, 4, 28), amount=D('75'),
        )
        self.assertEqual(self.line(self.build(2026, 5), e).other_deduction, D('75.00'))

    # ── أيام خارج فترة الاستحقاق ───────────────────────────────────────────
    def test_absence_before_hire_date_is_not_deducted(self):
        """مباشرة 16/4: غياب يوم 5/4 لا يُخصم (لا راتب أصلاً لذلك اليوم)."""
        e = self.emp(rate='0', hire=date(2026, 4, 16))
        EmployeeAbsence.objects.create(employee=e, absence_date=date(2026, 4, 5), days=1)
        self.assertEqual(self.line(self.build(2026, 4), e).absence_deduction, D('0.00'))

    # ── تغيير الراتب ───────────────────────────────────────────────────────
    def test_salary_change_after_draft_is_caught_before_lock(self):
        """زيادة الراتب بعد بناء المسودة: يجب أن يمنع التدقيقُ الترحيلَ بأرقام قديمة."""
        e = self.emp(rate='10')
        run = self.build(2026, 4)
        Employee.objects.filter(pk=e.pk).update(basic_salary=D('5000'))
        self.assertFalse(audit_payroll_runs([run]).ready_to_lock)

    def test_salary_change_after_draft_without_insurance_is_caught(self):
        e = self.emp(rate='0')
        run = self.build(2026, 4)
        Employee.objects.filter(pk=e.pk).update(transport_allowance=D('900'))
        self.assertFalse(audit_payroll_runs([run]).ready_to_lock)

    # ── إعادة فتح شهر سابق بعد ترحيل لاحق ──────────────────────────────────
    def test_cannot_unlock_month_when_later_month_is_locked(self):
        """فتح أبريل بعد ترحيل مايو يكسر تسلسل أرصدة المخصصات — يجب أن يُمنع."""
        self.emp(rate='0', hire=date(2023, 1, 1))
        apr = self.build(2026, 4)
        lock_payroll_run(apr, self.user)
        lock_payroll_run(self.build(2026, 5), self.user)
        with self.assertRaises(ValueError):
            unlock_payroll_run(apr, self.user)

    # ── رصيد الإجازة: إجازة سنوية مأخوذة داخل الشهر ─────────────────────────
    def test_leave_taken_inside_month_is_kept_in_cumulative_balance(self):
        """رصيد 10 أيام، أُخذت 5 أيام في 10/4، ثم رُحِّل أبريل (+1.75) ⇒ 6.75 لا 11.75."""
        e = self.emp(rate='0', hire=date(2023, 1, 1))
        EmployeeLedger.objects.create(
            employee=e, transaction_type=EmployeeLedger.TransactionType.INITIAL_BALANCE,
            date=date(2026, 3, 31), cumulative_leave_days=D('10'),
            cumulative_leave_amount=D('1500'), cumulative_eosb_amount=D('1000'),
        )
        EmployeeLedger.objects.create(
            employee=e, transaction_type=EmployeeLedger.TransactionType.LEAVE_TAKEN,
            date=date(2026, 4, 10), leave_days_change=D('-5'), leave_amount_change=D('-750'),
            cumulative_leave_days=D('5'), cumulative_leave_amount=D('750'),
            cumulative_eosb_amount=D('1000'),
        )
        run = self.build(2026, 4)
        lock_payroll_run(run, self.user)
        lg = EmployeeLedger.objects.get(payroll_run=run, employee=e)
        self.assertEqual(lg.cumulative_leave_days, D('6.75'))

    # ── التصفية بعد ترحيل الشهر ───────────────────────────────────────────
    def test_settlement_does_not_repay_salary_already_paid_in_locked_payroll(self):
        """أبريل رُحِّل كاملاً ثم صُفّي الموظف بتاريخ 20/4: راتب الفترة في التصفية يجب ألا يُصرف ثانيةً."""
        from apps.employees.services.settlement_financials import compute_settlement_financials
        e = self.emp(rate='0')
        lock_payroll_run(self.build(2026, 4), self.user)
        fin = compute_settlement_financials(e, date(2026, 4, 20))
        self.assertEqual(fin['prorated_salary'], D('0.00'))

    def test_settlement_deducts_pending_penalties_and_cash_shortages(self):
        """جزاءات وعجز كاشير لم تُحتسب في مسير بعد يجب أن تُخصم من التصفية."""
        from apps.employees.services.settlement_financials import compute_settlement_financials
        e = self.emp(rate='0')
        EmployeeStatement.objects.create(
            employee=e, statement_type=EmployeeStatement.StatementType.PENALTY,
            title='مخالفة', statement_date=date(2026, 4, 5), deduction_amount=D('300'),
        )
        EmployeeCashShortage.objects.create(
            employee=e, branch=self.branch, shortage_date=date(2026, 4, 6), amount=D('120'),
        )
        fin = compute_settlement_financials(e, date(2026, 4, 20))
        self.assertEqual(fin['total_deductions'], D('420.00'))

    # ── موظف حُذف (سلة المحذوفات) بعد بناء المسودة ─────────────────────────
    def test_employee_deleted_after_draft_is_not_paid_on_lock(self):
        e = self.emp(rate='0')
        keep = self.emp('باقٍ', rate='0')
        run = self.build(2026, 4)
        e.delete()  # حذف ناعم
        try:
            lock_payroll_run(run, self.user)
        except ValueError:
            return  # رفض الترحيل مقبول
        self.assertFalse(EmployeeLedger.objects.filter(payroll_run=run, employee_id=e.pk).exists())
        self.assertTrue(EmployeeLedger.objects.filter(payroll_run=run, employee=keep).exists())

    # ── تغيّر نوع الصرف بعد الترحيل ────────────────────────────────────────
    def test_mode_switch_after_lock_does_not_pay_twice(self):
        """رُحِّل في مسير التحويل ثم أُزيلت كفالته: لا يدخل مسير النقدي لنفس الشهر."""
        e = self.emp(rate='0')
        lock_payroll_run(self.build(2026, 4), self.user)
        Employee.objects.filter(pk=e.pk).update(sponsorship=None)
        cash = self.build(2026, 4, sponsored=False)
        self.assertFalse(cash.lines.filter(employee=e).exists())


# ══════════════════════════════════════════════════════════════════════════════
# 9) تعديل الراتب بتاريخ سريان — يسري على شهر تاريخه كاملاً
# ══════════════════════════════════════════════════════════════════════════════
class SalaryChangeEffectiveMonthTests(FinanceBase):
    def _raise(self, e, new_basic, eff, as_of):
        from apps.employees.services.salary_changes import record_salary_change
        return record_salary_change(e, new_basic=D(new_basic), effective_date=eff,
                                    user=self.user, as_of=as_of)

    def test_raise_dated_mid_month_applies_to_whole_month(self):
        """زيادة بتاريخ 20/10 ⇒ مسير أكتوبر كله بالأساسي الجديد (لا تقسيم)."""
        e = self.emp(rate='0')
        c = self._raise(e, '4000', date(2026, 10, 20), as_of=date(2026, 10, 8))
        self.assertTrue(c.applied)
        ln = self.line(self.build(2026, 10), e)
        self.assertEqual(ln.basic_salary, D('4000.00'))
        self.assertEqual(ln.gross_salary, D('5500.00'))

    def test_raise_from_next_month_does_not_touch_current_month(self):
        """زيادة من 1/11 سُجّلت في أكتوبر ⇒ أكتوبر بالقديم، نوفمبر بالجديد، والملف لا يتغير الآن."""
        e = self.emp(rate='0')
        c = self._raise(e, '4000', date(2026, 11, 1), as_of=date(2026, 10, 8))
        self.assertFalse(c.applied)
        e.refresh_from_db()
        self.assertEqual(e.basic_salary, D('3000.00'))
        self.assertEqual(self.line(self.build(2026, 10), e).basic_salary, D('3000.00'))
        self.assertEqual(self.line(self.build(2026, 11), e).basic_salary, D('4000.00'))

    def test_pending_raise_is_applied_when_its_month_arrives(self):
        from apps.employees.services.salary_changes import apply_due_salary_changes
        e = self.emp(rate='0')
        self._raise(e, '4000', date(2026, 11, 1), as_of=date(2026, 10, 8))
        self.assertEqual(apply_due_salary_changes(as_of=date(2026, 10, 31)), 0)
        self.assertEqual(apply_due_salary_changes(as_of=date(2026, 11, 1)), 1)
        e.refresh_from_db()
        self.assertEqual(e.basic_salary, D('4000.00'))

    def test_earlier_month_built_after_raise_uses_old_salary_and_audit_is_clean(self):
        e = self.emp(rate='10')
        self._raise(e, '4000', date(2026, 10, 1), as_of=date(2026, 10, 8))
        sep = self.build(2026, 9)
        self.assertEqual(self.line(sep, e).basic_salary, D('3000.00'))
        audit = audit_payroll_runs([sep])
        self.assertEqual(audit.error_count, 0, [c.detail for c in audit.checks if c.level == 'error'])

    def test_retro_raise_for_locked_month_is_paid_as_arrears_once(self):
        """سبتمبر رُحِّل بالقديم، ثم زيادة بتاريخ 15/9 ⇒ أكتوبر = جديد + فروقات سبتمبر، ونوفمبر بلا فروقات."""
        e = self.emp(rate='10')
        lock_payroll_run(self.build(2026, 9), self.user)
        self._raise(e, '4000', date(2026, 9, 15), as_of=date(2026, 10, 8))
        octo = self.build(2026, 10)
        ln = self.line(octo, e)
        self.assertEqual(ln.basic_salary, D('4000.00'))
        self.assertEqual(ln.other_addition, D('1000.00'))                 # فروقات سبتمبر
        self.assertEqual(ln.total_earnings, ln.gross_salary + D('1000.00'))
        self.assertEqual(ln.insurance_deduction, D('600.00'))             # 10% × (5000 + 1000)
        self.assertEqual(ln.net_salary, ln.total_earnings - ln.total_deductions)
        self.assertEqual(audit_payroll_runs([octo]).error_count, 0)
        lock_payroll_run(octo, self.user)
        self.assertEqual(self.line(self.build(2026, 11), e).other_addition, D('0.00'))

    def test_arrears_return_if_paying_month_is_reopened(self):
        e = self.emp(rate='0')
        lock_payroll_run(self.build(2026, 9), self.user)
        self._raise(e, '4000', date(2026, 9, 1), as_of=date(2026, 10, 8))
        octo = self.build(2026, 10)
        lock_payroll_run(octo, self.user)
        unlock_payroll_run(octo, self.user)
        self.assertEqual(self.line(self.build(2026, 10), e).other_addition, D('1000.00'))

    def test_arrears_are_prorated_for_partial_month(self):
        """موظف بدأ 16/9 (15 يوماً) ⇒ فروقات سبتمبر = 1000 × 15 ÷ 30 = 500."""
        e = self.emp(rate='0', hire=date(2026, 9, 16))
        lock_payroll_run(self.build(2026, 9), self.user)
        self._raise(e, '4000', date(2026, 9, 16), as_of=date(2026, 10, 8))
        self.assertEqual(self.line(self.build(2026, 10), e).other_addition, D('500.00'))

    def test_pending_action_with_future_date_does_not_change_salary_now(self):
        from types import SimpleNamespace
        from unittest import mock
        from apps.core.services.pending_actions import _execute_salary_adjust
        e = self.emp(rate='0')
        action = SimpleNamespace(
            employee=e, requested_by=self.user,
            payload={'new_basic_salary': '4500', 'effective_date': '2026-11-01', 'reason': 'ترقية'},
        )
        with mock.patch('django.utils.timezone.localdate', return_value=date(2026, 10, 8)):
            msg = _execute_salary_adjust(action, self.user)
        e.refresh_from_db()
        self.assertEqual(e.basic_salary, D('3000.00'))
        self.assertIn('تلقائياً', msg)
        self.assertEqual(self.line(self.build(2026, 11), e).basic_salary, D('4500.00'))


class FullExportWorkbookTests(FinanceBase):
    def test_single_export_has_all_sheets_and_consistent_totals(self):
        from io import BytesIO

        from django.urls import reverse
        from openpyxl import load_workbook

        sp2 = Sponsorship.objects.create(code='SP2', company_name='كفالة ثانية')
        self.emp('ت1', branch=self.branch)
        self.emp('ت2', branch=self.branch2)
        e3 = self.emp('ت3', branch=self.branch2)
        Employee.objects.filter(pk=e3.pk).update(sponsorship=sp2)
        self.emp('ن1', branch=self.branch, sponsored=False)

        for br in (self.branch, self.branch2):
            for sp in (self.sp, sp2):
                build_payroll_run(br, 2026, 4, self.user, salary_mode=TRANSFER, sponsorship_id=sp.id)
            build_payroll_run(br, 2026, 4, self.user, salary_mode=CASH)
        runs = PayrollRun.objects.filter(period_year=2026, period_month=4)
        expected_net = float(sum(r.total_net for r in runs))

        admin = User.objects.create_superuser(username='exp_admin', password='x-pass-123')
        self.client.force_login(admin)
        resp = self.client.get(reverse('web:export_payroll_list_excel'),
                               {'year': 2026, 'month': 4, 'salary_mode': 'transfer'})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('payroll_2026_04_full', resp['Content-Disposition'])
        wb = load_workbook(BytesIO(resp.content))
        self.assertEqual(wb.sheetnames,
                         ['كشف الرواتب', 'التحليل', 'حسب الشركات', 'التحويل', 'النقدي', 'التفصيلي'])

        # كشف الرواتب يضم الكل (3 تحويل + 1 نقدي) + صف إجمالي
        self.assertEqual(wb['كشف الرواتب'].max_row, 1 + 4 + 1)
        self.assertEqual(wb['التحويل'].max_row, 1 + 3 + 1)
        self.assertEqual(wb['النقدي'].max_row, 1 + 1 + 1)

        # التحليل: صف الإجمالي (عمود الصافي = 6) يساوي مجموع صافي المسيرات
        ws = wb['التحليل']
        total_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'الإجمالي')
        self.assertAlmostEqual(ws.cell(total_row, 6).value, expected_net, places=2)
        self.assertEqual(ws.cell(total_row, 2).value, 4)

        # حسب الشركات: جدول لكل شركة + جدول النقدي
        titles = [ws_c.value for ws_c in wb['حسب الشركات']['A'] if ws_c.value and '—' in str(ws_c.value)]
        self.assertTrue(any(t.startswith('كفالة الفحص') for t in titles), titles)
        self.assertTrue(any(t.startswith('كفالة ثانية') for t in titles), titles)
        self.assertTrue(any(t.startswith('نقدي') for t in titles), titles)
