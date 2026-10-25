"""استيراد/تصدير ربط البصمة بالموظفين (Excel)."""
import io
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.attendance.models import (
    AttendancePunch,
    BiometricDevice,
    BiometricEnrollmentAuditLog,
    EmployeeBiometricEnrollment,
)
from apps.core.models import Branch, Company, Permission, Role, UserProfile
from apps.employees.models import Employee

User = get_user_model()

HEADERS = ['الجهاز', 'رقم المستخدم على الجهاز', 'الاسم على الجهاز', 'الرقم الوظيفي', 'رقم الهوية']


def _xlsx(rows, headers=HEADERS, name='links.xlsx'):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(
        name, buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@override_settings(ALLOWED_HOSTS=['testserver'])
class EnrollmentExcelTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        company = Company.objects.create(name='شركة')
        cls.branch = Branch.objects.create(name='فرع', code='BR-EX', company=company)
        cls.other_branch = Branch.objects.create(name='فرع ثاني', code='BR-EY', company=company)
        cls.device = BiometricDevice.objects.create(name='جهاز 1', ip_address='10.0.0.1', branch=cls.branch)
        cls.emp1 = Employee.objects.create(
            name='علي', employee_number='E1', id_number='1000000001', branch=cls.branch,
            basic_salary=Decimal('3000'),
        )
        cls.emp2 = Employee.objects.create(
            name='سعيد', employee_number='E2', id_number='1000000002', branch=cls.branch,
            basic_salary=Decimal('3000'),
        )
        cls.admin = User.objects.create_user(
            username='en_admin', password='x12345678', is_superuser=True, is_staff=True,
        )

    def _import(self, upload, user=None):
        self.client.force_login(user or self.admin)
        return self.client.post(
            reverse('web:biometric_enrollments_import'), {'excel_file': upload}, follow=True,
        )

    def _messages(self, response):
        return [str(m) for m in response.context['messages']]

    # ── التصدير ──
    def test_export_contains_links(self):
        EmployeeBiometricEnrollment.objects.create(
            employee=self.emp1, device=self.device, device_user_id=7, device_user_name='Ali',
        )
        self.client.force_login(self.admin)
        r = self.client.get(reverse('web:biometric_enrollments_export'))
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r['Content-Type'])
        rows = list(load_workbook(io.BytesIO(r.content)).active.values)
        self.assertEqual(list(rows[0][:5]), HEADERS)
        self.assertEqual(list(rows[1][:5]), ['جهاز 1', 7, 'Ali', 'E1', '1000000001'])

    def test_export_requires_permission(self):
        user = User.objects.create_user(username='nobody', password='x12345678')
        self.client.force_login(user)
        r = self.client.get(reverse('web:biometric_enrollments_export'))
        self.assertNotEqual(r.status_code, 200)

    # ── الاستيراد ──
    def test_import_creates_links_and_relinks_punches(self):
        from django.utils import timezone
        AttendancePunch.objects.create(
            device=self.device, device_user_id=5, punched_at=timezone.now(),
        )
        r = self._import(_xlsx([['جهاز 1', 5, 'Ali', 'E1', None], ['جهاز 1', 6, '', None, '1000000002']]))
        self.assertEqual(EmployeeBiometricEnrollment.objects.filter(is_deleted=False).count(), 2)
        link = EmployeeBiometricEnrollment.objects.get(device_user_id=5)
        self.assertEqual(link.employee, self.emp1)
        self.assertEqual(AttendancePunch.objects.get(device_user_id=5).employee, self.emp1)
        self.assertEqual(EmployeeBiometricEnrollment.objects.get(device_user_id=6).employee, self.emp2)
        self.assertEqual(BiometricEnrollmentAuditLog.objects.filter(action='create').count(), 2)
        self.assertTrue(any('ربط جديد 2' in m for m in self._messages(r)))

    def test_reimport_is_idempotent(self):
        rows = [['جهاز 1', 5, 'Ali', 'E1', None]]
        self._import(_xlsx(rows))
        r = self._import(_xlsx(rows))
        self.assertEqual(EmployeeBiometricEnrollment.objects.filter(is_deleted=False).count(), 1)
        self.assertEqual(BiometricEnrollmentAuditLog.objects.count(), 1)
        self.assertTrue(any('بدون تغيير' in m for m in self._messages(r)))

    def test_changing_name_updates(self):
        self._import(_xlsx([['جهاز 1', 5, 'Ali', 'E1', None]]))
        self._import(_xlsx([['جهاز 1', 5, 'Ali Hassan', 'E1', None]]))
        self.assertEqual(EmployeeBiometricEnrollment.objects.get(device_user_id=5).device_user_name, 'Ali Hassan')
        self.assertEqual(BiometricEnrollmentAuditLog.objects.filter(action='update').count(), 1)

    def test_number_linked_to_other_employee_is_not_moved(self):
        EmployeeBiometricEnrollment.objects.create(employee=self.emp1, device=self.device, device_user_id=5)
        r = self._import(_xlsx([['جهاز 1', 5, '', 'E2', None]]))
        self.assertEqual(EmployeeBiometricEnrollment.objects.get(device_user_id=5).employee, self.emp1)
        self.assertTrue(any('مربوط بالموظف' in m for m in self._messages(r)))

    def test_employee_own_number_can_change(self):
        EmployeeBiometricEnrollment.objects.create(employee=self.emp1, device=self.device, device_user_id=5)
        self._import(_xlsx([['جهاز 1', 9, '', 'E1', None]]))
        links = EmployeeBiometricEnrollment.objects.filter(employee=self.emp1, is_deleted=False)
        self.assertEqual([l.device_user_id for l in links], [9])

    def test_bad_rows_reported_good_rows_saved(self):
        r = self._import(_xlsx([
            ['جهاز 1', 5, '', 'E1', None],
            ['جهاز غير موجود', 6, '', 'E2', None],
            ['جهاز 1', 'abc', '', 'E2', None],
            ['جهاز 1', 8, '', 'NOPE', None],
        ]))
        self.assertEqual(EmployeeBiometricEnrollment.objects.filter(is_deleted=False).count(), 1)
        text = ' '.join(self._messages(r))
        self.assertIn('تعذّر استيراد 3', text)
        self.assertIn('جهاز غير موجود', text)

    def test_duplicate_number_in_file_rejected(self):
        self._import(_xlsx([['جهاز 1', 5, '', 'E1', None], ['جهاز 1', 5, '', 'E2', None]]))
        self.assertEqual(EmployeeBiometricEnrollment.objects.get(device_user_id=5).employee, self.emp1)

    def test_single_device_may_omit_device_column(self):
        self._import(_xlsx([[5, 'Ali', 'E1']], headers=['رقم المستخدم على الجهاز', 'الاسم على الجهاز', 'الرقم الوظيفي']))
        self.assertEqual(EmployeeBiometricEnrollment.objects.get(device_user_id=5).device, self.device)

    def test_missing_columns_and_wrong_file(self):
        r = self._import(_xlsx([[1]], headers=['شيء']))
        self.assertTrue(any('رقم المستخدم' in m for m in self._messages(r)))
        r = self._import(SimpleUploadedFile('x.txt', b'hello'))
        self.assertTrue(any('xlsx' in m for m in self._messages(r)))
        self.assertFalse(EmployeeBiometricEnrollment.objects.exists())

    # ── الصلاحيات والنطاق ──
    def test_import_requires_manage_permission(self):
        role = Role.objects.create(name='مشاهد بصمة', role_type=Role.RoleType.HR_OFFICER)
        role.permissions.add(*Permission.objects.filter(code__in=['attendance_screen_devices.view']))
        user = User.objects.create_user(username='viewer', password='x12345678')
        profile, _ = UserProfile.objects.get_or_create(user=user)
        profile.role = role
        profile.save(update_fields=['role'])
        user = User.objects.select_related('profile__role').get(pk=user.pk)
        self._import(_xlsx([['جهاز 1', 5, '', 'E1', None]]), user=user)
        self.assertFalse(EmployeeBiometricEnrollment.objects.exists())

    def test_employee_outside_branch_scope_rejected(self):
        from apps.attendance.services.enrollment_excel import import_enrollments_from_upload
        summary = import_enrollments_from_upload(
            _xlsx([['جهاز 1', 5, '', 'E1', None]]), request=None,
            allowed_branch_ids={self.other_branch.id},
        )
        self.assertEqual(summary.errors, 1)
        self.assertFalse(EmployeeBiometricEnrollment.objects.exists())

    def test_buttons_shown_on_both_pages(self):
        self.client.force_login(self.admin)
        for name in ('web:attendance_daily', 'web:biometric_devices'):
            html = self.client.get(reverse(name)).content.decode()
            self.assertIn(reverse('web:biometric_enrollments_export'), html, name)
            self.assertIn(reverse('web:biometric_enrollments_import'), html, name)
