"""فلترة قائمة الموظفين: فرع، قسم، شركة كفالة، تاريخ المباشرة — في القائمة والتصدير."""
from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from apps.core.models import Branch, Company
from apps.core.selectors.employee_filters import (
    employee_list_querystring,
    parse_employee_list_filters,
)
from apps.departments.models import Department
from apps.employees.models import Employee
from apps.setup.models import Sponsorship

User = get_user_model()


@override_settings(ALLOWED_HOSTS=['testserver'])
class EmployeeListFilterTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        company = Company.objects.create(name='شركة')
        cls.b1 = Branch.objects.create(name='فرع أول', code='F1', company=company)
        cls.b2 = Branch.objects.create(name='فرع ثان', code='F2', company=company)
        cls.d1 = Department.objects.create(code='DD1', name='قسم أ', branch=cls.b1)
        cls.sp = Sponsorship.objects.create(code='SPF', company_name='كفالة الفلترة', is_active=True)
        Employee.objects.create(
            name='أحمد الأول', branch=cls.b1, department=cls.d1, sponsorship=cls.sp,
            hire_date=date(2025, 1, 10), status=Employee.Status.ACTIVE,
        )
        Employee.objects.create(
            name='سالم الثاني', branch=cls.b2, hire_date=date(2026, 3, 5),
            status=Employee.Status.ACTIVE,
        )
        cls.admin = User.objects.create_user(
            username='flt_admin', password='x12345678', is_superuser=True, is_staff=True,
        )

    def setUp(self):
        self.client = Client()
        self.client.force_login(self.admin)

    def _names(self, **params):
        response = self.client.get(reverse('web:list_employees'), params)
        self.assertEqual(response.status_code, 200)
        return {e.name for e in response.context['employees']}

    def test_parse_and_querystring(self):
        filters = parse_employee_list_filters({'branch': '3', 'from': '2026-01-01', 'to': 'bad'})
        self.assertEqual(filters['branch_id'], 3)
        self.assertIsNone(filters['date_to'])
        self.assertEqual(
            employee_list_querystring('x', filters), 'q=x&branch=3&from=2026-01-01',
        )

    def test_filters_narrow_results(self):
        self.assertEqual(self._names(), {'أحمد الأول', 'سالم الثاني'})
        self.assertEqual(self._names(branch=self.b2.pk), {'سالم الثاني'})
        self.assertEqual(self._names(department=self.d1.pk), {'أحمد الأول'})
        self.assertEqual(self._names(sponsorship=self.sp.pk), {'أحمد الأول'})
        self.assertEqual(self._names(**{'from': '2026-01-01'}), {'سالم الثاني'})
        self.assertEqual(self._names(**{'to': '2025-12-31'}), {'أحمد الأول'})

    def test_page_shows_filter_box(self):
        response = self.client.get(reverse('web:list_employees'))
        self.assertContains(response, 'الفلترة')
        self.assertContains(response, 'شركة الكفالة')

    def test_exports_respect_filters(self):
        url = reverse('web:export_sponsored_employees_excel')
        response = self.client.get(url, {'branch': self.b2.pk})
        rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual(len(rows), 1)  # العناوين فقط: لا موظف كفالة في الفرع الثاني
        response = self.client.get(url, {'branch': self.b1.pk})
        rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual(len(rows), 2)
