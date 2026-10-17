"""استيراد موظفين بكفالة من Excel (بيانات الموظف + الراتب + البنك).

الأعمدة المدعومة (تُطابَق بالاسم بعد تطبيع الهمزات والمسافات):
رقم الموظف، الاسم، رقم الآيبان، البنك، طبيعة الحساب، رقم الهوية، الفرع، الشركة،
الراتب الأساسي، بدل السكن، بدل النقل، بدل إضافي ثابت / حوافز، بدل إضافي، الجنس،
الجنسية، القسم، مركز التكلفة، تاريخ المباشرة.

يعيد استخدام أدوات الأمان والأداء من non_sponsored_import (حد الحجم، التحقق من xlsx،
كاش المراجع، الحفظ على دفعات).
"""
from __future__ import annotations

import logging
import re
from datetime import date
from decimal import Decimal
from typing import Any

from django.db import transaction
from django.db.models import Q
from django.utils import timezone

from apps.employees.models import Employee, SalaryAccountType
from apps.employees.services.non_sponsored_import import (
    IMPORT_SAVE_BATCH_SIZE,
    MAX_IMPORT_ROWS,
    ImportRowResult,
    ImportSummary,
    ImportValidationError,
    _FIELD_MAX_LEN,
    _ReferenceCache,
    _cell_date,
    _cell_decimal,
    _cell_str,
    _index_lookup,
    _resolve_from_index,
    _row_dict,
    _sanitize_text,
    validate_excel_upload,
)
from apps.setup.models import Bank, Sponsorship

logger = logging.getLogger(__name__)

_ARABIC_DIACRITICS = re.compile(r'[ً-ٰٟـ]')


def normalize_header(value: Any) -> str:
    """تطبيع عنوان العمود: همزات، تاء مربوطة، مسافات، فواصل."""
    text = '' if value is None else str(value)
    text = _ARABIC_DIACRITICS.sub('', text)
    for src, dst in (('أ', 'ا'), ('إ', 'ا'), ('آ', 'ا'), ('ى', 'ي'), ('ة', 'ه')):
        text = text.replace(src, dst)
    text = re.sub(r'\s+', ' ', text.replace('\n', ' ')).strip().lower()
    return text


# عناوين الأعمدة (مطبَّعة) → حقل الاستيراد
_HEADER_ALIASES = {
    'رقم الموظف': 'employee_number',
    'الرقم الوظيفي': 'employee_number',
    'الاسم': 'name',
    'اسم الموظف': 'name',
    'رقم الايبان': 'iban',
    'الايبان': 'iban',
    'iban': 'iban',
    'البنك': 'bank',
    'طبيعه الحساب': 'account_type',
    'رقم الهويه': 'id_number',
    'رقم الهويه / الاقامه': 'id_number',
    'الفرع': 'branch',
    'الشركه': 'sponsorship',
    'الكفاله': 'sponsorship',
    'الراتب الاساسي': 'basic_salary',
    'بدل السكن': 'housing_allowance',
    'بدل النقل': 'transport_allowance',
    'بدل اضافي': 'other_allowance',
    'بدل اضافي ثابت / حوافز': 'incentive_allowance',
    'الجنس': 'gender',
    'الجنسيه': 'nationality',
    'القسم': 'department',
    'مركز التكلفه': 'cost_center',
    'تاريخ المباشره': 'hire_date',
    'رقم الجوال': 'phone',
    'الجوال': 'phone',
    'المهنه': 'profession',
}

_SALARY_FIELDS = (
    'basic_salary', 'housing_allowance', 'transport_allowance',
    'other_allowance', 'incentive_allowance',
)

_ACCOUNT_TYPE_ALIASES = {
    'bank account': SalaryAccountType.BANK_ACCOUNT,
    'bank_account': SalaryAccountType.BANK_ACCOUNT,
    'حساب بنكي': SalaryAccountType.BANK_ACCOUNT,
    'salary card': SalaryAccountType.SALARY_CARD,
    'salary_card': SalaryAccountType.SALARY_CARD,
    'بطاقه راتب': SalaryAccountType.SALARY_CARD,
    'sarie': SalaryAccountType.SARIE,
    'سريع': SalaryAccountType.SARIE,
}

_GENDER_ALIASES = {
    'ذكر': Employee.Gender.MALE, 'male': Employee.Gender.MALE, 'm': Employee.Gender.MALE,
    'انثي': Employee.Gender.FEMALE, 'انثى': Employee.Gender.FEMALE,
    'female': Employee.Gender.FEMALE, 'f': Employee.Gender.FEMALE,
}


def field_for_header(label: Any) -> str | None:
    norm = normalize_header(label)
    if not norm:
        return None
    field = _HEADER_ALIASES.get(norm)
    if field:
        return field
    # «بدل إضافي ثاب / حوافز» وأي صيغة قريبة فيها «حوافز»
    if 'حوافز' in norm:
        return 'incentive_allowance'
    return None


def _parse_header_map(header_row) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        field = field_for_header(cell)
        if field and field not in mapping:
            mapping[field] = idx
    return mapping


def _normalize_iban(value: str) -> str:
    return re.sub(r'\s+', '', value or '').upper()[:34]


def _money(raw: dict, header_map: dict[str, int], key: str) -> Decimal | None:
    if key not in header_map:
        return None
    return _cell_decimal(raw.get(key))


def _load_existing(employee_numbers: set[str], id_numbers: set[str]):
    by_number: dict[str, Employee] = {}
    by_id: dict[str, Employee] = {}
    if not employee_numbers and not id_numbers:
        return by_number, by_id
    q = Q()
    if employee_numbers:
        q |= Q(employee_number__in=employee_numbers)
    if id_numbers:
        q |= Q(id_number__in=id_numbers)
    for emp in Employee.objects.filter(is_deleted=False).filter(q):
        if emp.employee_number:
            by_number.setdefault(emp.employee_number, emp)
        if emp.id_number:
            by_id.setdefault(emp.id_number, emp)
    return by_number, by_id


_UPDATE_FIELDS = [
    'name', 'sponsorship', 'employee_number', 'id_number', 'phone', 'hire_date',
    'branch', 'nationality', 'profession', 'department', 'cost_center', 'gender',
    'bank', 'iban', 'account_type',
    'basic_salary', 'housing_allowance', 'transport_allowance',
    'other_allowance', 'incentive_allowance', 'updated_at',
]


def import_sponsored_employees_from_workbook(
    workbook,
    *,
    user,
    allowed_branch_ids: set[int] | None,
    apply_salary: bool = True,
) -> ImportSummary:
    """ينشئ أو يحدّث موظفين بكفالة من الشيت الأول."""
    del user
    summary = ImportSummary()
    ws = workbook.active
    rows_iter = ws.iter_rows(values_only=True)

    try:
        header_row = next(rows_iter)
    except StopIteration:
        summary.add_result(ImportRowResult(1, 'error', 'الملف فارغ'))
        return summary

    header_map = _parse_header_map(header_row)
    if 'name' not in header_map:
        summary.add_result(ImportRowResult(1, 'error', 'عمود «الاسم» مطلوب في الصف الأول'))
        return summary

    raw_rows: list[tuple[int, tuple]] = []
    for excel_row_num, row in enumerate(rows_iter, start=2):
        if row is None or all(c is None or str(c).strip() == '' for c in row):
            continue
        raw_rows.append((excel_row_num, row))
        if len(raw_rows) > MAX_IMPORT_ROWS:
            raise ImportValidationError(
                f'عدد الصفوف يتجاوز الحد المسموح ({MAX_IMPORT_ROWS}). قسّم الملف إلى أجزاء.',
            )

    refs = _ReferenceCache(allowed_branch_ids)
    sponsorship_index: dict[str, Sponsorship] = {}
    for sp in Sponsorship.objects.filter(is_deleted=False, is_active=True):
        for key in ((sp.company_name or '').strip().lower(), (sp.code or '').strip().lower()):
            if key and key not in sponsorship_index:
                sponsorship_index[key] = sp
    bank_index = _index_lookup(Bank.objects.filter(is_deleted=False, is_active=True))

    employee_numbers: set[str] = set()
    id_numbers: set[str] = set()
    for _num, row in raw_rows:
        raw = _row_dict(row, header_map)
        en = _sanitize_text(_cell_str(raw.get('employee_number')), _FIELD_MAX_LEN['employee_number'])
        idn = _sanitize_text(_cell_str(raw.get('id_number')), _FIELD_MAX_LEN['id_number'])
        if en:
            employee_numbers.add(en)
        if idn:
            id_numbers.add(idn)
    by_number, by_id = _load_existing(employee_numbers, id_numbers)

    pending_creates: list[tuple[Employee, ImportRowResult]] = []
    pending_updates: list[tuple[Employee, ImportRowResult]] = []

    def _remember(emp: Employee) -> None:
        if emp.employee_number:
            by_number[emp.employee_number] = emp
        if emp.id_number:
            by_id[emp.id_number] = emp

    def flush_pending() -> None:
        nonlocal pending_creates, pending_updates
        if not pending_creates and not pending_updates:
            return
        with transaction.atomic():
            for emp, result in pending_creates:
                emp.save()
                _remember(emp)
                summary.add_result(result)
            for emp, result in pending_updates:
                emp.save(update_fields=_UPDATE_FIELDS)
                _remember(emp)
                summary.add_result(result)
        pending_creates = []
        pending_updates = []

    def _ref(index, header_key: str, raw: dict, label: str):
        if header_key not in header_map:
            return None
        text = _cell_str(raw.get(header_key))
        if not text:
            return None
        obj = _resolve_from_index(index, text)
        if obj is None:
            raise ImportValidationError(f'{label} غير موجود: {text}')
        return obj

    for excel_row_num, row in raw_rows:
        raw = _row_dict(row, header_map)
        name = _sanitize_text(_cell_str(raw.get('name')), _FIELD_MAX_LEN['name'])
        if not name:
            summary.add_result(ImportRowResult(excel_row_num, 'skipped', 'اسم فارغ — تم التخطي'))
            continue
        employee_number = _sanitize_text(
            _cell_str(raw.get('employee_number')), _FIELD_MAX_LEN['employee_number'],
        )
        try:
            id_number = _sanitize_text(_cell_str(raw.get('id_number')), _FIELD_MAX_LEN['id_number'])
            phone = _sanitize_text(_cell_str(raw.get('phone')), _FIELD_MAX_LEN['phone'])
            hire_date: date | None = (
                _cell_date(raw.get('hire_date')) if 'hire_date' in header_map else None
            )

            branch = None
            if 'branch' in header_map:
                branch_raw = _cell_str(raw.get('branch'))
                if branch_raw:
                    branch = _resolve_from_index(refs.branches, branch_raw)
                    if branch is None:
                        raise ImportValidationError(
                            f'الفرع غير موجود أو خارج صلاحيتك: {branch_raw}',
                        )

            sponsorship = None
            if 'sponsorship' in header_map:
                sp_raw = _cell_str(raw.get('sponsorship'))
                if sp_raw:
                    sponsorship = sponsorship_index.get(sp_raw.strip().lower())
                    if sponsorship is None:
                        raise ImportValidationError(f'الشركة (الكفالة) غير موجودة: {sp_raw}')

            nationality = _ref(refs.nationalities, 'nationality', raw, 'الجنسية')
            profession = _ref(refs.professions, 'profession', raw, 'المهنة')
            department = _ref(refs.departments, 'department', raw, 'القسم')
            cost_center = _ref(refs.cost_centers, 'cost_center', raw, 'مركز التكلفة')
            bank = _ref(bank_index, 'bank', raw, 'البنك')

            gender = ''
            if 'gender' in header_map:
                g_raw = _cell_str(raw.get('gender'))
                if g_raw:
                    gender = _GENDER_ALIASES.get(normalize_header(g_raw), '')
                    if not gender:
                        raise ImportValidationError(f'الجنس غير معروف: {g_raw}')

            account_type = ''
            if 'account_type' in header_map:
                a_raw = _cell_str(raw.get('account_type'))
                if a_raw:
                    account_type = _ACCOUNT_TYPE_ALIASES.get(normalize_header(a_raw), '')
                    if not account_type:
                        raise ImportValidationError(f'طبيعة الحساب غير معروفة: {a_raw}')

            iban = ''
            if 'iban' in header_map:
                iban = _normalize_iban(_cell_str(raw.get('iban')))

            amounts = {}
            if apply_salary:
                for key in _SALARY_FIELDS:
                    value = _money(raw, header_map, key)
                    if value is not None:
                        amounts[key] = value

            existing = None
            if employee_number:
                existing = by_number.get(employee_number)
            if existing is None and id_number:
                existing = by_id.get(id_number)
            if existing is not None and allowed_branch_ids is not None \
                    and existing.branch_id not in allowed_branch_ids:
                raise ImportValidationError(
                    f'الموظف موجود في فرع خارج صلاحيتك ({existing.employee_number or existing.id_number})',
                )

            is_create = existing is None
            if is_create:
                if sponsorship is None:
                    raise ImportValidationError('الشركة (الكفالة) مطلوبة لموظف جديد بكفالة')
                emp = Employee(status=Employee.Status.ACTIVE)
            else:
                emp = existing

            emp.name = name
            if sponsorship is not None:
                emp.sponsorship = sponsorship
            if employee_number:
                emp.employee_number = employee_number
            if id_number:
                emp.id_number = id_number
            if phone:
                emp.phone = phone
            if hire_date is not None:
                emp.hire_date = hire_date
            if branch is not None:
                emp.branch = branch
            elif is_create and allowed_branch_ids is not None and len(allowed_branch_ids) == 1:
                emp.branch_id = next(iter(allowed_branch_ids))
            if is_create and emp.branch_id is None and allowed_branch_ids is not None:
                raise ImportValidationError('الفرع مطلوب عند إنشاء موظف جديد')
            if nationality is not None:
                emp.nationality = nationality
            if profession is not None:
                emp.profession = profession
            if department is not None:
                emp.department = department
            if cost_center is not None:
                emp.cost_center = cost_center
            if gender:
                emp.gender = gender
            if bank is not None:
                emp.bank = bank
            if iban:
                emp.iban = iban
            if account_type:
                emp.account_type = account_type
            for key, value in amounts.items():
                setattr(emp, key, value)
            emp.updated_at = timezone.now()

            result = ImportRowResult(
                excel_row_num,
                'created' if is_create else 'updated',
                f'تم {"الإنشاء" if is_create else "التحديث"}: {emp.name}',
                employee_number=emp.employee_number or '',
            )
            (pending_creates if is_create else pending_updates).append((emp, result))
            if len(pending_creates) + len(pending_updates) >= IMPORT_SAVE_BATCH_SIZE:
                flush_pending()

        except ImportValidationError as exc:
            summary.add_result(
                ImportRowResult(excel_row_num, 'error', str(exc), employee_number=employee_number),
            )
        except Exception:
            logger.exception('sponsored import row %s failed', excel_row_num)
            summary.add_result(
                ImportRowResult(
                    excel_row_num, 'error', 'تعذّر حفظ الصف لأسباب داخلية.',
                    employee_number=employee_number,
                ),
            )

    flush_pending()
    return summary


def import_sponsored_employees_from_upload(
    uploaded_file,
    *,
    user,
    allowed_branch_ids,
    apply_salary=True,
) -> ImportSummary:
    from openpyxl import load_workbook

    validate_excel_upload(uploaded_file)
    wb = load_workbook(uploaded_file, read_only=True, data_only=True)
    try:
        return import_sponsored_employees_from_workbook(
            wb,
            user=user,
            allowed_branch_ids=allowed_branch_ids,
            apply_salary=apply_salary,
        )
    finally:
        wb.close()
