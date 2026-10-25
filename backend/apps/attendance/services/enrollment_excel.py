"""تصدير واستيراد ربط البصمة بالموظفين (Excel).

الأعمدة (نفس ترتيب التصدير ويقبلها الاستيراد):
    الجهاز | رقم المستخدم على الجهاز | الاسم على الجهاز | الرقم الوظيفي | رقم الهوية | اسم الموظف

الاستيراد يحدّد الموظف بـ «الرقم الوظيفي» ثم «رقم الهوية» (واحد منهما يكفي)،
والجهاز باسمه (يمكن حذف العمود إن وُجد جهاز واحد فقط).
قواعد الأمان: لا يُنقل رقم مستخدم مربوط بموظف آخر (يُتخطّى ويُبلَّغ عنه)،
وكل تغيير يُسجَّل في سجل تدقيق الربط نفسه الذي تستعمله شاشة الأجهزة.
"""
from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from typing import Any

from django.db import transaction

from apps.attendance.models import (
    AttendancePunch,
    BiometricDevice,
    BiometricDeviceUser,
    BiometricEnrollmentAuditLog,
    EmployeeBiometricEnrollment,
)
from apps.attendance.services.ingest_audit import log_enrollment_change
from apps.employees.models import Employee
from apps.employees.services.non_sponsored_import import (
    ImportValidationError,
    _cell_str,
    _norm_header,
    validate_excel_upload,
)

logger = logging.getLogger(__name__)

MAX_ROWS = 5000
MAX_DETAIL_RESULTS = 40

COLUMNS = [
    ('الجهاز', 'device'),
    ('رقم المستخدم على الجهاز', 'device_user_id'),
    ('الاسم على الجهاز', 'device_user_name'),
    ('الرقم الوظيفي', 'employee_number'),
    ('رقم الهوية', 'id_number'),
    ('اسم الموظف', 'employee_name'),
]

HEADER_ALIASES = {
    'الجهاز': 'device',
    'اسم الجهاز': 'device',
    'رقم المستخدم على الجهاز': 'device_user_id',
    'رقم المستخدم': 'device_user_id',
    'رقم البصمة': 'device_user_id',
    'الاسم على الجهاز': 'device_user_name',
    'الرقم الوظيفي': 'employee_number',
    'رقم الموظف': 'employee_number',
    'رقم الهوية': 'id_number',
    'اسم الموظف': 'employee_name',  # للعرض فقط — لا يُستعمل في المطابقة
}


@dataclass
class RowResult:
    row_number: int
    action: str  # created | updated | unchanged | skipped | error
    message: str


@dataclass
class ImportSummary:
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    skipped: int = 0
    errors: int = 0
    results: list[RowResult] = field(default_factory=list)

    _COUNTERS = {
        'created': 'created', 'updated': 'updated', 'unchanged': 'unchanged',
        'skipped': 'skipped', 'error': 'errors',
    }

    def add(self, row_number: int, action: str, message: str) -> None:
        counter = self._COUNTERS[action]
        setattr(self, counter, getattr(self, counter) + 1)
        if action in ('skipped', 'error') and len(self.results) < MAX_DETAIL_RESULTS:
            self.results.append(RowResult(row_number, action, message))


# ───────────────────────────── التصدير ─────────────────────────────

def build_enrollments_workbook_bytes(*, allowed_branch_ids=None) -> bytes:
    """ملف Excel بكل روابط البصمة النشطة (يُقيَّد بفروع المستخدم عند وجود نطاق)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    qs = (
        EmployeeBiometricEnrollment.objects.filter(is_deleted=False, employee__is_deleted=False)
        .select_related('employee', 'device')
        .order_by('device__name', 'device_user_id')
    )
    if allowed_branch_ids is not None:
        qs = qs.filter(employee__branch_id__in=allowed_branch_ids)

    wb = Workbook()
    ws = wb.active
    ws.title = 'ربط البصمة'
    ws.sheet_view.rightToLeft = True
    header_fill = PatternFill('solid', fgColor='2B1B5E')
    for col, (label, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for row, e in enumerate(qs.iterator(chunk_size=500), start=2):
        values = [
            e.device.name, e.device_user_id, e.device_user_name,
            e.employee.employee_number, e.employee.id_number, e.employee.name,
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
    for col, width in enumerate((26, 22, 28, 18, 18, 34), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ───────────────────────────── الاستيراد ─────────────────────────────

def _parse_headers(header_row) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, value in enumerate(header_row):
        key = HEADER_ALIASES.get(_norm_header(value))
        if key and key not in mapping:
            mapping[key] = idx
    return mapping


def _get(row, mapping: dict[str, int], key: str) -> Any:
    idx = mapping.get(key)
    return row[idx] if idx is not None and idx < len(row) else None


def _resolve_device(name: str, devices_by_name: dict[str, list], only_device):
    if not name:
        if only_device is None:
            raise ValueError('اسم الجهاز مطلوب (يوجد أكثر من جهاز).')
        return only_device
    matches = devices_by_name.get(name.casefold(), [])
    if not matches:
        raise ValueError(f'جهاز غير موجود: «{name}».')
    if len(matches) > 1:
        raise ValueError(f'يوجد أكثر من جهاز بالاسم «{name}» — غيّر الاسم ليصبح فريداً.')
    return matches[0]


def _resolve_employee(number: str, id_number: str, by_number: dict, by_id: dict):
    if number:
        found = by_number.get(number)
        if found is None:
            raise ValueError(f'لا يوجد موظف بالرقم الوظيفي «{number}».')
        if len(found) > 1:
            raise ValueError(f'الرقم الوظيفي «{number}» مكرر عند أكثر من موظف.')
        return found[0]
    if id_number:
        found = by_id.get(id_number)
        if found is None:
            raise ValueError(f'لا يوجد موظف برقم الهوية «{id_number}».')
        if len(found) > 1:
            raise ValueError(f'رقم الهوية «{id_number}» مكرر عند أكثر من موظف.')
        return found[0]
    raise ValueError('الرقم الوظيفي أو رقم الهوية مطلوب.')


def import_enrollments_from_workbook(wb, *, request, allowed_branch_ids=None) -> ImportSummary:
    summary = ImportSummary()
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return summary
    mapping = _parse_headers(header)
    if 'device_user_id' not in mapping:
        raise ImportValidationError('العمود «رقم المستخدم على الجهاز» غير موجود في الملف.')
    if 'employee_number' not in mapping and 'id_number' not in mapping:
        raise ImportValidationError('يلزم عمود «الرقم الوظيفي» أو «رقم الهوية».')

    devices = list(BiometricDevice.objects.filter(is_deleted=False))
    devices_by_name: dict[str, list] = {}
    for d in devices:
        devices_by_name.setdefault((d.name or '').strip().casefold(), []).append(d)
    only_device = devices[0] if len(devices) == 1 else None
    if 'device' not in mapping and only_device is None:
        raise ImportValidationError('يوجد أكثر من جهاز — أضف عمود «الجهاز».')

    employees = Employee.objects.filter(is_deleted=False).only(
        'id', 'name', 'employee_number', 'id_number', 'branch_id',
    )
    by_number: dict[str, list] = {}
    by_id: dict[str, list] = {}
    for emp in employees:
        if emp.employee_number:
            by_number.setdefault(emp.employee_number.strip(), []).append(emp)
        if emp.id_number:
            by_id.setdefault(emp.id_number.strip(), []).append(emp)

    seen_users: set[tuple[int, int]] = set()
    seen_employees: set[tuple[int, int]] = set()

    for offset, row in enumerate(rows_iter, start=2):
        if offset - 1 > MAX_ROWS:
            raise ImportValidationError(f'عدد الصفوف يتجاوز الحد المسموح ({MAX_ROWS}).')
        if row is None or all(c is None or str(c).strip() == '' for c in row):
            continue
        try:
            device = _resolve_device(_cell_str(_get(row, mapping, 'device')), devices_by_name, only_device)

            raw_id = _cell_str(_get(row, mapping, 'device_user_id'))
            if not raw_id.isdigit() or int(raw_id) <= 0:
                raise ValueError(f'رقم المستخدم على الجهاز غير صالح: «{raw_id}».')
            device_user_id = int(raw_id)

            employee = _resolve_employee(
                _cell_str(_get(row, mapping, 'employee_number')),
                _cell_str(_get(row, mapping, 'id_number')),
                by_number, by_id,
            )
            if allowed_branch_ids is not None and employee.branch_id not in allowed_branch_ids:
                raise ValueError(f'الموظف «{employee.name}» خارج الفروع المتاحة لك.')

            if (device.pk, device_user_id) in seen_users:
                raise ValueError(f'الرقم {device_user_id} مكرر في الملف لنفس الجهاز.')
            if (device.pk, employee.pk) in seen_employees:
                raise ValueError(f'الموظف «{employee.name}» مكرر في الملف لنفس الجهاز.')
            seen_users.add((device.pk, device_user_id))
            seen_employees.add((device.pk, employee.pk))

            name = _cell_str(_get(row, mapping, 'device_user_name'))[:200]
            if not name:
                name = (
                    BiometricDeviceUser.objects.filter(device=device, device_user_id=device_user_id)
                    .values_list('name', flat=True).first()
                ) or ''

            action, message = _apply_link(request, device, device_user_id, employee, name)
            summary.add(offset, action, message)
        except ValueError as exc:
            summary.add(offset, 'error', str(exc))
        except Exception:
            logger.exception('biometric enrollment import row %s failed', offset)
            summary.add(offset, 'error', 'تعذّر حفظ هذا الصف.')
    return summary


def _apply_link(request, device, device_user_id: int, employee, name: str) -> tuple[str, str]:
    """ينشئ/يحدّث الربط بنفس قواعد شاشة الأجهزة. يُرجع (الإجراء، الرسالة)."""
    active = EmployeeBiometricEnrollment.objects.filter(device=device, is_deleted=False)
    by_user = active.filter(device_user_id=device_user_id).select_related('employee').first()
    if by_user is not None and by_user.employee_id != employee.pk:
        raise ValueError(
            f'الرقم {device_user_id} مربوط بالموظف «{by_user.employee.name}» — لم يُغيَّر. '
            'أعد ربطه من شاشة الأجهزة إن كان مقصوداً.'
        )
    by_employee = active.filter(employee=employee).first()

    if by_user is not None and (by_user.device_user_name or '') == name:
        return 'unchanged', ''
    if by_user is None and by_employee is None:
        audit_action = BiometricEnrollmentAuditLog.Action.CREATE
        result = 'created'
    else:
        audit_action = BiometricEnrollmentAuditLog.Action.UPDATE
        result = 'updated'

    with transaction.atomic():
        # الموظف مربوط على الجهاز برقم آخر ← يُستبدل (يتعارض مع قيد جهاز+موظف)
        EmployeeBiometricEnrollment.all_objects.filter(
            device=device, employee=employee,
        ).exclude(device_user_id=device_user_id).hard_delete()
        EmployeeBiometricEnrollment.all_objects.update_or_create(
            device=device, device_user_id=device_user_id,
            defaults={
                'employee': employee, 'device_user_name': name,
                'is_deleted': False, 'deleted_at': None,
            },
        )
        relinked = AttendancePunch.objects.filter(
            device=device, device_user_id=device_user_id, is_deleted=False,
        ).update(employee=employee, device_user_name=name)
    log_enrollment_change(
        request=request, device=device, device_user_id=device_user_id,
        new_employee=employee, device_user_name=name,
        action=audit_action, punches_relinked=relinked,
    )
    return result, ''


def import_enrollments_from_upload(uploaded_file, *, request, allowed_branch_ids=None) -> ImportSummary:
    from openpyxl import load_workbook

    validate_excel_upload(uploaded_file)
    wb = load_workbook(uploaded_file, read_only=True, data_only=True)
    try:
        return import_enrollments_from_workbook(
            wb, request=request, allowed_branch_ids=allowed_branch_ids,
        )
    finally:
        wb.close()
